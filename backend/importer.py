import openpyxl
import docx
import os
import sqlite3
import datetime
import re
import csv

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "gastro_roster.db")

def parse_date(val):
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, str):
        val = val.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%b-%y", "%d-%m-%Y"):
            try:
                return datetime.datetime.strptime(val[:10], fmt).strftime("%Y-%m-%d")
            except:
                pass
    return None

CUSTOM_MAP = {
    'Alex Thompson': 'Prof. Alex Thompson',
    'Prof Alex Thompson': 'Prof. Alex Thompson',
    'Jacinta Holmes': 'Assoc Prof Jacinta Holmes',
    'Prof Jessica Howell': 'Assoc Prof Jessica Howell',
    'Jessica Howell': 'Assoc Prof Jessica Howell',
    'Dan Tassone': 'Dan Tassone',
    'Daniel Tassone': 'Dan Tassone',
    'Nik Vootukuru': 'Nikil Vootukuru',
    'Nikil Vootukuru': 'Nikil Vootukuru',
    'Kyle Wu': 'Dr Kyle Wu',
    'Declan Connoley': 'Dr Declan Connoley',
    'Tara Fox': 'Dr Tara Fox',
    'John Ding': 'Assoc Prof John Nik Ding',
    'Emily Wright': 'Assoc Prof Emily Wright',
    'Ashley Miller': 'Assoc Prof Ashley Miller',
    'Bill Connell': 'Assoc Prof Bill Connell',
    'Paul Desmond': 'Assoc Prof Paul Desmond',
    'Marno Ryan': 'Assoc Prof Marno Ryan',
    'Andrew Taylor': 'Assoc Prof Andrew Taylor',
    'Bronte Holt': 'Assoc Prof Bronte Holt',
    'Simon Hume': 'Dr Simon Hume',
    'Tamie Samyue': 'Tamie Samyue',
    'Hannah Kim': 'Dr Hannah Kim',
    'Chamara Basnayake': 'Dr Chamara Basnayake',
    'Georgie Cameron': 'Dr Georgie Cameron',
    'Catherine Croagh': 'Dr Catherine Croagh',
    'David Iser': 'Dr David Iser',
    'Edward Tsoi': 'Dr Edward Tsoi',
    'Emma Flanagan': 'Dr Emma Flanagan',
    'Linda Yang': 'Dr Linda Yang',
    'Mark Lust': 'Dr Mark Lust',
    'Michael De Gregorio': 'Dr Michael De Gregorio',
    'Ola Niewiadomski': 'Dr Ola Niewiadomski',
    'Ralley Prentice': 'Dr Ralley Prentice',
    'Robert Chen': 'Dr Robert Chen',
    'Sasha Fehily': 'Dr Sasha Fehily',
    'Suzanne Mahady': 'Dr Suzanne Mahady',
    'Swee Lin Chen Yi Mei': 'Dr Swee Lin Chen Yi Mei',
    'Thai Hong': 'Dr Thai Hong',
    'Tim Papaluca': 'Dr Tim Papaluca',
    'Tony He': 'Dr Tony He',
    'Ann Farrell': 'Dr Ann Farrell',
    'Barbara Demediuk': 'Dr Barbara Demediuk',
    'Imogen Hartley': 'Dr Imogen Hartley',
    'Tanya Lee': 'Dr Tanya Lee',
    'Lauren Andersson': 'Dr Lauren Andersson',
    'Nick Clark': 'Dr Nick Clark',
    'Kathryn Gazelakis': 'Dr Kathryn Gazelakis'
}

CLINIC_DOCTOR_MAP = {
    'ALEX': 'Prof. Alex Thompson',
    'ANDREW': 'Assoc Prof Andrew Taylor',
    'ANN': 'Dr Ann Farrell',
    'ASHLEY': 'Assoc Prof Ashley Miller',
    'BARB': 'Dr Barbara Demediuk',
    'BILL': 'Assoc Prof Bill Connell',
    'CATHERINE': 'Dr Catherine Croagh',
    'CHAMARA': 'Dr Chamara Basnayake',
    'DAVID': 'Dr David Iser',
    'ED': 'Dr Edward Tsoi',
    'EMILY': 'Assoc Prof Emily Wright',
    'EMMA': 'Dr Emma Flanagan',
    'FELLOW': 'RESEARCH FELLOWS',
    'GEORGIE': 'Dr Georgie Cameron',
    'HANNAH (IBD PSYCH)': 'Dr Hannah Kim',
    'JACINTA': 'Assoc Prof Jacinta Holmes',
    'JESS': 'Assoc Prof Jessica Howell',
    'JOHN': 'Assoc Prof John Nik Ding',
    'LINDA': 'Dr Linda Yang',
    'MARK': 'Dr Mark Lust',
    'MARNO': 'Assoc Prof Marno Ryan',
    'OLA': 'Dr Ola Niewiadomski',
    'PAUL': 'Assoc Prof Paul Desmond',
    'RALLEY': 'Dr Ralley Prentice',
    'REG': 'Registrar',
    'SASHA': 'Dr Sasha Fehily',
    'SIMON': 'Dr Simon Hume',
    'SUZANNE': 'Dr Suzanne Mahady',
    'SWEELIN': 'Dr Swee Lin Chen Yi Mei',
    'TAMIE': 'Tamie Samyue',
    'THAI': 'Dr Thai Hong',
    'TIM': 'Dr Tim Papaluca',
    'TONY': 'Dr Tony He'
}

def clean_name(name):
    if not name:
        return ""
    name = str(name).strip()
    return CUSTOM_MAP.get(name, name)


def import_all():
    from database import init_db
    init_db()
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    
    # Clear existing entries
    cur.execute("DELETE FROM roster_entries")
    cur.execute("DELETE FROM staff")

    staff_set = set()

    # 1. Staff roster from SMS Sheet2
    sms_file = os.path.join(BASE_DIR, "SMS roster schedule Gastroenterology 2026 V2 270126.xlsx")
    if os.path.exists(sms_file):
        wb = openpyxl.load_workbook(sms_file, data_only=True)
        if "Sheet2" in wb.sheetnames:
            ws2 = wb["Sheet2"]
            for r in range(1, ws2.max_row + 1):
                raw = ws2.cell(r, 2).value
                if raw and str(raw).strip():
                    cleaned = clean_name(raw)
                    staff_set.add(cleaned)
                    is_coord = 1 if "Basnayake" in cleaned or "Thompson" in cleaned else 0
                    cur.execute(
                        "INSERT OR IGNORE INTO staff (name, email, role, is_coordinator) VALUES (?, ?, ?, ?)",
                        (cleaned, f"{cleaned.lower().replace(' ', '.').replace('dr.', '').replace('prof.', '').replace('assoc', '').replace('..', '.')[:15]}@svha.org.au", 'consultant', is_coord)
                    )

        # 2. Import SMS On-call, wards, consults from Sheet1
        ws1 = wb["Sheet1"]
        # Data rows start at 10
        for r in range(10, ws1.max_row + 1):
            date_val = parse_date(ws1.cell(r, 3).value)
            if not date_val:
                continue
            holiday = ws1.cell(r, 1).value
            holiday = str(holiday).strip() if holiday else None
            
            # HM Registrar (Col 12)
            reg = ws1.cell(r, 12).value
            if reg and str(reg).strip():
                s_name = clean_name(reg)
                cur.execute(
                    "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    ('hm_oncall', date_val, 'Monday-Sunday', 'H&M Oncall Registrar', s_name, 'Registrar', None, holiday, s_name)
                )

            # HM Consultant (Col 14)
            cons = ws1.cell(r, 14).value
            if cons and str(cons).strip():
                s_name = clean_name(cons)
                cur.execute(
                    "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    ('hm_oncall', date_val, 'Monday-Sunday', 'H&M Oncall Consultant', s_name, 'Consultant', None, holiday, s_name)
                )

            # IBD Wards (Col 17)
            ibd = ws1.cell(r, 17).value
            if ibd and str(ibd).strip():
                s_name = clean_name(ibd)
                cur.execute(
                    "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    ('ibd_wards', date_val, 'Monday-Sunday', 'IBD Ward Service', s_name, 'Consultant', None, holiday, s_name)
                )

            # Liver Wards (Col 19)
            liver = ws1.cell(r, 19).value
            if liver and str(liver).strip():
                s_name = clean_name(liver)
                cur.execute(
                    "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    ('liver_wards', date_val, 'Monday-Sunday', 'Liver Ward Service', s_name, 'Consultant', None, holiday, s_name)
                )

            # Consults (Col 21)
            consults = ws1.cell(r, 21).value
            if consults and str(consults).strip():
                s_name = clean_name(consults)
                cur.execute(
                    "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    ('consults', date_val, 'Monday-Sunday', 'Consults Consultant', s_name, 'Consultant', None, holiday, s_name)
                )

    # 3. Import GICC (Prefer GICC 2026.csv, fallback to GICC 2026 V2.xlsx)
    gicc_csv = os.path.join(BASE_DIR, "GICC 2026.csv")
    gicc_file = os.path.join(BASE_DIR, "GICC 2026 V2.xlsx")
    if os.path.exists(gicc_csv):
        with open(gicc_csv, "r", encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                if len(row) > 2 and row[2].strip() and (row[2].strip().startswith("2026-") or row[2].strip().startswith("2027-")):
                    date_val = row[2].strip()
                    holiday = row[0].strip() if row[0].strip() else None
                    meeting_format = row[3].strip() if len(row) > 3 and row[3].strip() else None
                    chair = row[4].strip() if len(row) > 4 else ""
                    spk1 = row[5].strip() if len(row) > 5 else ""
                    spk2 = row[6].strip() if len(row) > 6 else ""

                    if not chair and not spk1 and not spk2:
                        continue

                    if chair:
                        norm = clean_name(chair)
                        cur.execute(
                            "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                            ('gicc', date_val, 'Monday 12:30pm', 'GICC Meeting Chair', norm, 'Chair', str(meeting_format or ''), holiday, norm)
                        )
                    if spk1:
                        norm = clean_name(spk1)
                        cur.execute(
                            "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                            ('gicc', date_val, 'Monday 12:30pm', 'GICC Meeting Speaker 1', norm, 'Speaker 1', str(meeting_format or ''), holiday, norm)
                        )
                    if spk2:
                        norm = clean_name(spk2)
                        cur.execute(
                            "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                            ('gicc', date_val, 'Monday 12:30pm', 'GICC Meeting Speaker 2', norm, 'Speaker 2', str(meeting_format or ''), holiday, norm)
                        )
    elif os.path.exists(gicc_file):
        wb_g = openpyxl.load_workbook(gicc_file, data_only=True)
        if "GICC roster" in wb_g.sheetnames:
            ws_g = wb_g["GICC roster"]
            for r in range(10, ws_g.max_row + 1):
                date_val = parse_date(ws_g.cell(r, 3).value)
                if not date_val:
                    continue
                holiday = ws_g.cell(r, 1).value
                holiday = str(holiday).strip() if holiday else None
                meeting_format = ws_g.cell(r, 4).value # F2F or Zoom
                
                chair = ws_g.cell(r, 5).value
                if chair and str(chair).strip():
                    cur.execute(
                        "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        ('gicc', date_val, 'Monday 12:30pm', 'GICC Meeting Chair', clean_name(chair), 'Chair', str(meeting_format or ''), holiday, clean_name(chair))
                    )

                spk1 = ws_g.cell(r, 6).value
                if spk1 and str(spk1).strip():
                    cur.execute(
                        "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        ('gicc', date_val, 'Monday 12:30pm', 'GICC Meeting Speaker 1', clean_name(spk1), 'Speaker 1', str(meeting_format or ''), holiday, clean_name(spk1))
                    )

                spk2 = ws_g.cell(r, 7).value
                if spk2 and str(spk2).strip():
                    cur.execute(
                        "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        ('gicc', date_val, 'Monday 12:30pm', 'GICC Meeting Speaker 2', clean_name(spk2), 'Speaker 2', str(meeting_format or ''), holiday, clean_name(spk2))
                    )

    # 4. Import Friday journal club (Prefer Friday journal club 2026.csv, fallback to Friday journal club V1.xlsx)
    jc_csv = os.path.join(BASE_DIR, "Friday journal club 2026.csv")
    jc_file = os.path.join(BASE_DIR, "Friday journal club V1.xlsx")
    if os.path.exists(jc_csv):
        with open(jc_csv, "r", encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                if len(row) > 3 and row[2].strip() and (row[2].strip().startswith("2026-") or row[2].strip().startswith("2027-")):
                    date_val = row[2].strip()
                    presenter = row[3].strip()
                    if not presenter:
                        continue
                    holiday = row[0].strip() if row[0].strip() else None
                    norm = clean_name(presenter)
                    cur.execute(
                        "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        ('friday_jc', date_val, 'Friday 8:00am', 'Friday Journal Club', norm, 'Presenter', None, holiday, norm)
                    )
    elif os.path.exists(jc_file):
        wb_jc = openpyxl.load_workbook(jc_file, data_only=True)
        ws_jc = wb_jc.active
        for r in range(7, ws_jc.max_row + 1):
            date_val = parse_date(ws_jc.cell(r, 3).value)
            if not date_val:
                continue
            holiday = ws_jc.cell(r, 1).value
            holiday = str(holiday).strip() if holiday else None
            presenter = ws_jc.cell(r, 4).value
            if presenter and str(presenter).strip():
                cur.execute(
                    "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    ('friday_jc', date_val, 'Friday 8:00am', 'Friday Journal Club', clean_name(presenter), 'Presenter', None, holiday, clean_name(presenter))
                )

    # 5. Import 8.30 IBD meeting roster until Christmas.docx
    ibd_doc = os.path.join(BASE_DIR, "8.30 IBD meeting roster until Christmas.docx")
    if os.path.exists(ibd_doc):
        doc = docx.Document(ibd_doc)
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            # Format like '7/9: Review of HPOS paperwork' or '14/9: Journal Club. LY'
            match = re.match(r'^(\d{1,2})[/.](\d{1,2})[:.]?\s*(.*)', text)
            if match:
                day, month, details = match.groups()
                # Assuming 2026 for consistent calendar
                try:
                    d_obj = datetime.date(2026, int(month), int(day))
                    d_str = d_obj.strftime("%Y-%m-%d")
                    cur.execute(
                        "INSERT INTO roster_entries (entry_type, date, day_of_week, title, staff_name, role_label, notes, holiday, original_staff_name) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        ('ibd_meeting', d_str, 'Monday 8:30am', 'IBD Morning Meeting', details, 'Lead', details, None, details)
                    )
                except Exception as e:
                    pass

    # 6. Import Clinic Leave Roster (2026-2028)
    import_clinic_roster(conn)

    # 7. Import Functional Gut Clinic (Thursdays 2026-2028)
    import_functional_gut_clinic(conn)

    # 8. Import SMS Monthly Leave Update
    import_sms_leave(conn)

    # 9. Import Theatre Grid (Endoscopy Lists)
    import_theatre_grid(conn)

    conn.commit()
    cur.execute("SELECT count(*) FROM roster_entries")
    count = cur.fetchone()[0]
    cur.execute("SELECT count(*) FROM staff")
    staff_count = cur.fetchone()[0]
    cur.execute("SELECT count(*) FROM clinic_roster")
    clinic_count = cur.fetchone()[0]
    cur.execute("SELECT count(*) FROM theatre_grid_entries")
    theatre_count = cur.fetchone()[0]
    conn.close()
    print(f"Import complete: {count} roster entries, {staff_count} staff, {clinic_count} clinic entries, {theatre_count} theatre entries.")

VIC_PUBLIC_HOLIDAYS_2026 = {
    "2026-01-01": "PUBLIC HOLIDAY — New Year's Day (no clinic)",
    "2026-01-26": "PUBLIC HOLIDAY — Australia Day (no clinic)",
    "2026-03-09": "PUBLIC HOLIDAY — Labour Day (no clinic)",
    "2026-04-03": "PUBLIC HOLIDAY — Good Friday (no clinic)",
    "2026-04-06": "PUBLIC HOLIDAY — Easter Monday (no clinic)",
    "2026-04-25": "PUBLIC HOLIDAY — ANZAC Day (no clinic)",
    "2026-06-08": "PUBLIC HOLIDAY — King's Birthday (no clinic)",
    "2026-09-25": "PUBLIC HOLIDAY — Friday before the AFL Grand Final (no clinic)",
    "2026-11-03": "PUBLIC HOLIDAY — Melbourne Cup Day (no clinic)",
    "2026-12-25": "PUBLIC HOLIDAY — Christmas Day (no clinic)",
    "2026-12-28": "PUBLIC HOLIDAY — Boxing Day (substitute) (no clinic)",
}

def import_clinic_roster(conn):
    cur = conn.cursor()
    cur.execute("DELETE FROM clinic_roster")
    
    # Ensure Simon Hume and Tamie Samyue in staff
    cur.execute("INSERT OR IGNORE INTO staff (name, email, role, is_coordinator) VALUES (?, ?, ?, ?)",
                ("Dr Simon Hume", "simon.hume@svha.org.au", "consultant", 0))
    cur.execute("INSERT OR IGNORE INTO staff (name, email, role, is_coordinator) VALUES (?, ?, ?, ?)",
                ("Tamie Samyue", "tamie.samyue@svha.org.au", "fellow", 0))

    clinic_file = os.path.join(BASE_DIR, "StVs_Clinic_Leave_Roster_2026-2028.xlsx")
    if not os.path.exists(clinic_file):
        print("Clinic roster file not found:", clinic_file)
        return

    wb = openpyxl.load_workbook(clinic_file, data_only=True)
    sheets = ['GASTRO', 'OES', 'RAE MONDAY', 'RAE TUESDAY', 'RAE FRIDAY', 'WEIGHT MANAGEMENT', 'LIVER', 'IBD', 'PO LIVER', 'HEPATOMA']

    for sheet_name in sheets:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = [str(ws.cell(1, c).value or '').strip() for c in range(1, ws.max_column + 1)]

        # 1. Parse Excel rows and collect clean non-holiday template rows for backfill
        parsed_excel_rows = []
        clean_template_rows = []
        session_weekdays = set()

        for r in range(2, ws.max_row + 1):
            d_raw = ws.cell(r, 1).value
            d_val = parse_date(d_raw)
            if not d_val:
                continue
            day_of_week = str(ws.cell(r, 2).value or '').strip()
            session_weekdays.add(day_of_week)

            row_str = ' '.join([str(ws.cell(r, c).value or '') for c in range(1, ws.max_column + 1)])
            is_holiday_row = 'PUBLIC HOLIDAY' in row_str.upper()
            holiday_note = None
            if is_holiday_row:
                for c in range(1, ws.max_column + 1):
                    val_c = str(ws.cell(r, c).value or '')
                    if 'PUBLIC HOLIDAY' in val_c.upper():
                        holiday_note = val_c.strip()
                        break

            parsed_excel_rows.append((d_val, day_of_week, r, is_holiday_row, holiday_note))
            if not is_holiday_row and 'CLOSED' not in row_str.upper():
                clean_template_rows.append((d_val, day_of_week, r))

        if not parsed_excel_rows:
            continue

        first_excel_dt_str = parsed_excel_rows[0][0]
        first_excel_dt = datetime.datetime.strptime(first_excel_dt_str, "%Y-%m-%d").date()

        # 2. Backfill from 2026-01-01 up to first_excel_dt - 1 day
        backfill_curr = datetime.date(2026, 1, 1)
        while backfill_curr < first_excel_dt:
            day_name = backfill_curr.strftime("%A")
            if day_name in session_weekdays:
                d_str = backfill_curr.strftime("%Y-%m-%d")
                is_vic_holiday = d_str in VIC_PUBLIC_HOLIDAYS_2026
                vic_holiday_note = VIC_PUBLIC_HOLIDAYS_2026.get(d_str)

                # Find 4-week cycle template row from clean template rows
                matched_row_idx = None
                for t_dval, t_day, t_r in clean_template_rows:
                    if t_day == day_name:
                        t_dt = datetime.datetime.strptime(t_dval, "%Y-%m-%d").date()
                        if (t_dt - backfill_curr).days % 28 == 0:
                            matched_row_idx = t_r
                            break

                # Fallback to any clean row on same day of week if exact 28-day modulo not found
                if matched_row_idx is None:
                    for t_dval, t_day, t_r in clean_template_rows:
                        if t_day == day_name:
                            matched_row_idx = t_r
                            break

                for col_idx in range(3, len(headers) + 1):
                    doc_key = headers[col_idx - 1]
                    if not doc_key or doc_key in ['AMY', 'GEETHA']:
                        continue
                    staff_name = CLINIC_DOCTOR_MAP.get(doc_key)
                    if not staff_name:
                        continue

                    if is_vic_holiday:
                        cur.execute('''
                            INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        ''', (sheet_name, d_str, day_name, staff_name, 'holiday', None, 'not_required', vic_holiday_note))
                        continue

                    raw_cell = ws.cell(matched_row_idx, col_idx).value if matched_row_idx else None
                    cell_val = str(raw_cell).strip() if raw_cell is not None else ''
                    cell_upper = cell_val.upper()

                    if cell_val == '' or 'LEAVE' in cell_upper or cell_val in ['Unknown', 'jj']:
                        # Baseline scheduled attendance
                        cur.execute('''
                            INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        ''', (sheet_name, d_str, day_name, staff_name, 'attending', None, 'not_required', None))
                    elif cell_upper in ['X', '.']:
                        cur.execute('''
                            INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        ''', (sheet_name, d_str, day_name, staff_name, 'off_routine', None, 'not_required', None))
                    elif 'CLOSED' in cell_upper:
                        cur.execute('''
                            INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        ''', (sheet_name, d_str, day_name, staff_name, 'closed', None, 'not_required', cell_val))
                    else:
                        cur.execute('''
                            INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        ''', (sheet_name, d_str, day_name, staff_name, 'attending', None, 'not_required', cell_val))

            backfill_curr += datetime.timedelta(days=1)

        # 3. Insert original Excel rows (Dec 2026 onwards)
        for d_val, day_of_week, r, is_holiday_row, holiday_note in parsed_excel_rows:
            for col_idx in range(3, len(headers) + 1):
                doc_key = headers[col_idx - 1]
                if not doc_key or doc_key in ['AMY', 'GEETHA']:
                    continue
                staff_name = CLINIC_DOCTOR_MAP.get(doc_key)
                if not staff_name:
                    continue

                if is_holiday_row:
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (sheet_name, d_val, day_of_week, staff_name, 'holiday', None, 'not_required', holiday_note))
                    continue

                raw_cell = ws.cell(r, col_idx).value
                cell_val = str(raw_cell).strip() if raw_cell is not None else ''
                cell_upper = cell_val.upper()

                if cell_val == '':
                    # Attending scheduled clinic
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (sheet_name, d_val, day_of_week, staff_name, 'attending', None, 'not_required', None))
                elif cell_upper in ['X', '.']:
                    # Routine off
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (sheet_name, d_val, day_of_week, staff_name, 'off_routine', None, 'not_required', None))
                elif 'CLOSED' in cell_upper:
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (sheet_name, d_val, day_of_week, staff_name, 'closed', None, 'not_required', cell_val))
                elif 'LEAVE' in cell_upper or cell_val in ['Unknown', 'jj']:
                    # On leave -> Needs coverage!
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (sheet_name, d_val, day_of_week, staff_name, 'on_leave', cell_val, 'uncovered', None))
                else:
                    # Specialized clinic tag (e.g. HEP B, HEPATOMA) -> Attending with note
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (sheet_name, d_val, day_of_week, staff_name, 'attending', None, 'not_required', cell_val))

def import_functional_gut_clinic(conn):
    cur = conn.cursor()

    # 1. Ensure all Functional Gut Clinic clinicians exist in staff table
    fgc_staff = [
        ("Dr Chamara Basnayake", "chamara.basnayake@svha.org.au", "consultant", 1),
        ("Dr Ralley Prentice", "ralley.prentice@svha.org.au", "consultant", 0),
        ("Gastroenterology Registrar", "gastro.registrar@svha.org.au", "registrar", 0),
        ("Cherry De Crespigny", "cherry.decrespigny@svha.org.au", "allied_health", 0),
        ("Dr Hannah Kim", "hannah.kim@svha.org.au", "consultant", 0),
        ("Dr Jessica Green", "jessica.green@svha.org.au", "consultant", 0),
        ("Psychiatry Registrar", "psych.registrar@svha.org.au", "registrar", 0),
        ("Anna Stanley", "anna.stanley@svha.org.au", "allied_health", 0),
        ("Jacqueline Zaki", "jacqueline.zaki@svha.org.au", "allied_health", 0),
        ("Madeline Green", "madeline.green@svha.org.au", "allied_health", 0)
    ]

    for name, email, role, is_coord in fgc_staff:
        cur.execute(
            "INSERT OR IGNORE INTO staff (name, email, role, is_coordinator) VALUES (?, ?, ?, ?)",
            (name, email, role, is_coord)
        )

    # 2. Clear any prior Functional Gut clinic entries
    cur.execute("DELETE FROM clinic_roster WHERE clinic_name = 'FUNCTIONAL GUT'")

    # 3. Anchor: Thursday 15th October 2026 is Cycle A
    anchor_dt = datetime.date(2026, 10, 15)

    # Date range: 2026-01-01 to 2028-12-31
    start_dt = datetime.date(2026, 1, 1)
    end_dt = datetime.date(2028, 12, 31)

    public_holidays = {
        "2026-01-01": "PUBLIC HOLIDAY — New Year's Day (no clinic)",
        "2027-01-01": "PUBLIC HOLIDAY — New Year's Day (no clinic)",
        "2028-01-01": "PUBLIC HOLIDAY — New Year's Day (no clinic)"
    }

    curr = start_dt
    while curr <= end_dt:
        if curr.weekday() == 3: # Thursday
            d_str = curr.strftime("%Y-%m-%d")
            weeks_diff = (curr - anchor_dt).days // 7
            is_cycle_a = (weeks_diff % 2 == 0)

            is_holiday = d_str in public_holidays
            holiday_note = public_holidays.get(d_str)

            # Weekly team
            weekly_members = [
                ("Dr Chamara Basnayake", "Consultant (PM)"),
                ("Dr Ralley Prentice", "Consultant (PM)"),
                ("Gastroenterology Registrar", "Registrar (PM)"),
                ("Cherry De Crespigny", "Physiotherapist (PM)"),
                ("Dr Hannah Kim", "Psychiatrist (PM)"),
                ("Dr Jessica Green", "Psychiatrist (PM)"),
                ("Psychiatry Registrar", "Psychiatry Registrar (PM)")
            ]

            for s_name, role_note in weekly_members:
                if is_holiday:
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES ('FUNCTIONAL GUT', ?, 'Thursday', ?, 'holiday', NULL, 'not_required', ?)
                    ''', (d_str, s_name, holiday_note))
                else:
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES ('FUNCTIONAL GUT', ?, 'Thursday', ?, 'attending', NULL, 'not_required', ?)
                    ''', (d_str, s_name, role_note))

            # Alternating team:
            # Anna Stanley (Cycle A)
            # Jacqueline Zaki & Madeline Green (Cycle B)
            if is_holiday:
                for alt_name, role_label in [
                    ("Anna Stanley", "Dietitian"),
                    ("Jacqueline Zaki", "Physiotherapist"),
                    ("Madeline Green", "Dietitian")
                ]:
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES ('FUNCTIONAL GUT', ?, 'Thursday', ?, 'holiday', NULL, 'not_required', ?)
                    ''', (d_str, alt_name, holiday_note))
            else:
                if is_cycle_a:
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES ('FUNCTIONAL GUT', ?, 'Thursday', 'Anna Stanley', 'attending', NULL, 'not_required', 'Dietitian (Cycle A)')
                    ''', (d_str,))
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES ('FUNCTIONAL GUT', ?, 'Thursday', 'Jacqueline Zaki', 'off_routine', NULL, 'not_required', 'Physiotherapist (Off Cycle A)')
                    ''', (d_str,))
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES ('FUNCTIONAL GUT', ?, 'Thursday', 'Madeline Green', 'off_routine', NULL, 'not_required', 'Dietitian (Off Cycle A)')
                    ''', (d_str,))
                else:
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES ('FUNCTIONAL GUT', ?, 'Thursday', 'Anna Stanley', 'off_routine', NULL, 'not_required', 'Dietitian (Off Cycle B)')
                    ''', (d_str,))
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES ('FUNCTIONAL GUT', ?, 'Thursday', 'Jacqueline Zaki', 'attending', NULL, 'not_required', 'Physiotherapist (Cycle B)')
                    ''', (d_str,))
                    cur.execute('''
                        INSERT INTO clinic_roster (clinic_name, date, day_of_week, staff_name, status, leave_type, coverage_status, notes)
                        VALUES ('FUNCTIONAL GUT', ?, 'Thursday', 'Madeline Green', 'attending', NULL, 'not_required', 'Dietitian (Cycle B)')
                    ''', (d_str,))

        curr += datetime.timedelta(days=1)

def import_sms_leave(conn):
    cur = conn.cursor()
    sms_leave_file = os.path.join(BASE_DIR, "SMS monthly leave update V2(Dec 26).csv")
    if not os.path.exists(sms_leave_file):
        return

    with open(sms_leave_file, "r", encoding="latin1") as f:
        reader = csv.DictReader(f)
        for r in reader:
            raw_name = r.get('Name')
            if not raw_name or not str(raw_name).strip() or str(raw_name).strip() == 'Barb?':
                continue
            name = clean_name(str(raw_name).strip())
            start_d = parse_date(r.get('Leave start date'))
            end_d = parse_date(r.get('Leave end date'))
            if not start_d:
                continue
            if not end_d:
                end_d = start_d
            l_type = r.get('Type') or 'Annual'

            # Avoid exact duplicates
            cur.execute('''
                SELECT count(*) FROM leave_applications
                WHERE staff_name = ? AND start_date = ? AND end_date = ?
            ''', (name, start_d, end_d))
            if cur.fetchone()[0] == 0:
                cur.execute('''
                    INSERT INTO leave_applications (staff_name, leave_type, start_date, end_date, status, manager_name, submission_date, created_at)
                    VALUES (?, ?, ?, ?, 'approved', 'Department Coordinator', '2026-10-01', datetime('now'))
                ''', (name, l_type, start_d, end_d))

def import_theatre_grid(conn):
    import json
    cur = conn.cursor()
    cur.execute("DELETE FROM theatre_grid_entries")

    json_path = os.path.join(BASE_DIR, "frontend", "src", "theatre_grid_weeks.json")
    if not os.path.exists(json_path):
        return

    with open(json_path, "r") as f:
        weeks_data = json.load(f)

    # Anchor Monday: 2026-11-02 is Week 3
    anchor_dt = datetime.date(2026, 11, 2)
    anchor_week = 3

    # Generate cycle from 2026-10-05 (Monday) to 2027-12-27 (Monday)
    start_monday = datetime.date(2026, 10, 5)
    end_monday = datetime.date(2027, 12, 27)

    day_offsets = {
        "Monday": 0,
        "Tuesday": 1,
        "Wednesday": 2,
        "Thursday": 3,
        "Friday": 4
    }

    curr_monday = start_monday
    while curr_monday <= end_monday:
        weeks_diff = (curr_monday - anchor_dt).days // 7
        cycle_week = ((weeks_diff + (anchor_week - 1)) % 4) + 1
        week_key = str(cycle_week)

        if week_key in weeks_data and "sessions" in weeks_data[week_key]:
            for sess in weeks_data[week_key]["sessions"]:
                day_name = sess.get("day")
                offset = day_offsets.get(day_name, 0)
                sess_date = (curr_monday + datetime.timedelta(days=offset)).strftime("%Y-%m-%d")
                staff_name = clean_name(sess.get("staff_name") or "")
                room = sess.get("room") or "Endo"
                list_type = sess.get("listType") or "ENDO"
                session_time = sess.get("session") or "AM"
                notes = f"{sess.get('startTime', '')} - {sess.get('site', '')}".strip()

                cur.execute('''
                    INSERT INTO theatre_grid_entries (date, day_of_week, session, room_name, list_type, staff_name, coverage_status, notes)
                    VALUES (?, ?, ?, ?, ?, ?, 'not_required', ?)
                ''', (sess_date, day_name, session_time, room, list_type, staff_name, notes))

        curr_monday += datetime.timedelta(days=7)

if __name__ == "__main__":
    import_all()
